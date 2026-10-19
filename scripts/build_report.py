from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "AccessAble_AI_CA3_Final_Report.docx"

def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), color)
    props.append(node)

def paragraph(doc, text, center=False, mono=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    if mono:
        run.font.name = "Consolas"
    return p

def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, value in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = value
        shade(cell, "173F5F")
        for run in cell.paragraphs[0].runs:
            run.bold = True; run.font.color.rgb = RGBColor(255, 255, 255); run.font.size = Pt(9)
    for number, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            if number % 2:
                shade(cells[i], "EDF3F8")
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(9)
    doc.add_paragraph()

def h(doc, text, level=1):
    return doc.add_heading(text, level=level)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(.7); section.bottom_margin = Inches(.7)
section.left_margin = Inches(.8); section.right_margin = Inches(.8)
styles = doc.styles
styles["Normal"].font.name = "Aptos"; styles["Normal"].font.size = Pt(10.5)
for name in ("Title", "Heading 1", "Heading 2"):
    styles[name].font.name = "Aptos Display"; styles[name].font.color.rgb = RGBColor(0, 0, 0)

# Cover page
cover = doc.add_paragraph(); cover.paragraph_format.space_before = Pt(130); cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = cover.add_run("AccessAble AI"); run.bold = True; run.font.size = Pt(32); run.font.color.rgb = RGBColor(23, 63, 95)
paragraph(doc, "Multi Agent Website Accessibility Auditor", True).runs[0].font.size = Pt(18)
for item in ("CA 3 Agentic AI Domain Specific Proof of Concept", "Prepared by Priyansh Pandey", "Python and Streamlit", "October 2026"):
    paragraph(doc, item, True).runs[0].font.size = Pt(13)
doc.add_page_break()

sections = [
    ("Abstract", "AccessAble AI is a Streamlit proof of concept for web development and quality assurance. It performs a limited first-pass review of supplied HTML, then hands structured findings through deterministic agents for standards mapping, remediation guidance, validation boundaries, and completeness review. The broken synthetic store page produces five findings under the implemented rules, while the improved page produces zero. The application does not claim WCAG conformance because human review remains necessary."),
    ("Introduction and Domain Justification", "The selected domain is web development and quality assurance. Public and student-facing pages change quickly, while accessibility specialists may not review every early page. A focused tool can identify selected HTML barriers early and give developers a structured starting point for remediation."),
    ("Problem Statement", "Small development and QA teams can publish pages containing basic accessibility barriers without recognising them during routine implementation. Manual assessment remains necessary, but a reviewer benefits from a repeatable first-pass process that combines evidence, standards context, practical fixes, and a clear validation boundary."),
    ("Proposed Solution and Objectives", "The Streamlit interface accepts a synthetic demo page or pasted HTML. It runs five transparent rules, groups findings by severity, adds WCAG 2.2 context and implementation examples, offers CSV export, and displays a human-review checklist. Its objective is to demonstrate a testable agentic QA workflow rather than certify a website."),
    ("Technology Stack and Implementation", "The implementation uses Python 3.10, Streamlit, pandas, Python's built-in html.parser, and unittest. The contrast calculation uses relative luminance for supported inline colours. It avoids paid APIs and external LLM dependencies so the rule logic stays repeatable and inspectable."),
    ("Streamlit Interface and Demonstration Workflow", "A user chooses Broken store page, Improved store page, or Custom HTML and selects Run multi-agent audit. The interface displays counts, workflow, plan and critic output, expandable evidence, recommendations, human-review checks, and CSV download. Empty input receives a safe prompt and malformed input does not terminate the audit."),
    ("Output Validation and Human Review Requirements", "The Validation Agent states what each automatic check established. Human reviewers must judge meaningful image descriptions, keyboard operation and visible focus, screen-reader flow, dynamic updates, timing, motion, and cognitive clarity. The Critic checks report completeness only; it does not validate accessibility quality."),
    ("Limitations", "The scanner inspects supplied HTML only. It does not execute JavaScript, crawl a website, calculate computed styles, inspect all CSS inheritance or interaction states, test keyboard flow, or operate a screen reader. It cannot decide whether content is meaningful in context. It must not be used as a WCAG conformance certificate."),
    ("Future Scope", "Future work could add rendered-browser checks, broader CSS support, keyboard-test integration, batch audits, history, and a human reviewer approval workflow. A local language model could phrase recommendations, but deterministic checks and human review should remain the validation authority."),
    ("Conclusion", "AccessAble AI demonstrates agent orchestration for a focused QA problem. It converts selected HTML signals into an explainable developer report and explicitly preserves the point at which human review takes over."),
    ("Installation Run and Test Instructions", "From the project root, install dependencies with python -m pip install -r requirements.txt. Start the interface with streamlit run app.py. Run tests with python -m unittest discover -s tests -v. Use the broken page for the five-finding demonstration and the improved page for the zero-finding comparison."),
]
for title, text in sections[:4]:
    h(doc, title); paragraph(doc, text)

h(doc, "Existing Process and Identified Challenges")
table(doc, ["Challenge", "Effect"], [("Manual discovery of simple HTML omissions", "Reviewers spend time locating structural problems before analysing context."), ("Errors without standards context", "Developers need extra research to understand why a rule matters."), ("Overreliance on automation", "A zero-error result can be mistaken for full accessibility compliance.")])

h(doc, "System Architecture and Data Flow")
paragraph(doc, "The code implements the following structured hand-off. It has no external model, browser crawler, persistent memory, or autonomous browsing.")
paragraph(doc, "User / QA analyst → Streamlit UI → HTML input or sample page →\nPlanner → Scanner → Standards Analyst → Fix Recommendation →\nValidation → Critic → Severity report, CSV export and human checklist", True, True)
paragraph(doc, "Figure 1. Implemented AccessAble AI architecture and data flow.", True).runs[0].italic = True

h(doc, "Agent Roles Responsibilities and Hand Off Contracts")
table(doc, ["Stage", "Input", "Output and decision"], [("Planner Orchestrator", "HTML text", "Supported rule plan or safe no-input status."), ("Scanner Agent", "HTML and plan", "Selected structural signals as findings."), ("Standards Analyst", "Findings", "Relevant WCAG 2.2 criterion and rationale."), ("Fix Recommendation", "Analysed findings", "Implementation guidance and example."), ("Validation Agent", "Recommended findings", "Validation boundary and human checklist."), ("Critic Agent", "Validated findings", "Completeness notes for every finding.")])

h(doc, "Accessibility Rules and WCAG Mappings")
table(doc, ["Rule", "Automated signal", "WCAG 2.2 mapping"], [("IMG ALT", "Empty or absent alt", "1.1.1 Non-text Content"), ("FORM LABEL", "No linked label/simple name", "1.3.1 and 3.3.2"), ("PAGE LANG", "Missing html lang", "3.1.1 Language of Page"), ("PAGE TITLE", "No title text", "2.4.2 Page Titled"), ("COLOR CONTRAST", "Inline ratio under 4.5:1", "1.4.3 Contrast Minimum")])

for title, text in sections[4:]:
    h(doc, title); paragraph(doc, text)
    if title == "Streamlit Interface and Demonstration Workflow":
        h(doc, "Testing Methodology and Actual Results", 1)
        paragraph(doc, "The final verification ran seven unittest tests successfully using Python 3.10.11. The suite covers each sample result, contrast calculation, standards context, recommendations, validation guidance, empty input, malformed HTML, planner output, critic output, and checklist generation.")
        table(doc, ["Verification", "Actual result"], [("Automated tests", "7 passed"), ("Python syntax", "Compiled successfully"), ("Broken sample", "5 findings: 3 High, 2 Medium"), ("Improved sample", "0 findings under implemented rules"), ("Streamlit command", "Started locally during verification")])
        h(doc, "Before and After Findings Comparison", 1)
        table(doc, ["Demo page", "High", "Medium", "Total", "Interpretation"], [("Broken store page", "3", "2", "5", "Demonstrates the selected rules."), ("Improved store page", "0", "0", "0", "No finding under the five rules.")])

h(doc, "References")
for ref in ("W3C. Web Content Accessibility Guidelines WCAG 2.2. https://www.w3.org/TR/WCAG22/", "W3C. Easy Checks A First Review of Web Accessibility. https://www.w3.org/WAI/test-evaluate/preliminary/", "Streamlit Documentation. https://docs.streamlit.io/"):
    paragraph(doc, ref)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("AccessAble AI CA 3 Report")
OUT.parent.mkdir(parents=True, exist_ok=True); doc.save(OUT); print(OUT)
